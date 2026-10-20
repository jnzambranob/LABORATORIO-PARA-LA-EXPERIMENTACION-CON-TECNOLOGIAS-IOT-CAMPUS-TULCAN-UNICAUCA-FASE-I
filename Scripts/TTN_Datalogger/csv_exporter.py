"""
CSV and Excel Exporter module for TTN Data Collector.
Reads collected metrics from SQLite and generates formatted CSV and Excel (.xlsx) files
with full session-level isolation and multi-session comparison.
"""

import csv
from datetime import datetime, timezone
import logging
from pathlib import Path
from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db_storage import Database

logger = logging.getLogger("ttn_collector")


class DataExporter:
    """Exports SQLite metrics into CSV and multi-sheet Excel files with session isolation."""

    def __init__(self, db: Database, output_dir: Path):
        self.db = db
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_all(self, session_id: Optional[str] = None, prefix: Optional[str] = None) -> Dict[str, Path]:
        """
        Exports both CSV files and a comprehensive Excel file for a specific session (or latest).
        Returns a dictionary with paths to created files.
        """
        target_session = session_id or self.db.get_latest_session_id()
        now_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        sess_tag = target_session or now_str
        tag = f"{prefix}_{sess_tag}" if prefix else sess_tag

        csv_files = self.export_to_csv(tag=tag, session_id=target_session)
        excel_file = self.export_to_excel(tag=tag, session_id=target_session)

        results = {**csv_files, "excel": excel_file}
        logger.info(f"Exportación completada para sesión '{target_session}'. Archivos en: {self.output_dir}")
        return results

    def export_to_csv(self, tag: str, session_id: Optional[str] = None) -> Dict[str, Path]:
        """Generates individual CSV files for uplinks, downlinks, RTT analysis, and gateway stats."""
        created_files = {}

        # 1. Uplinks CSV
        uplinks = self.db.fetch_all_uplinks(session_id=session_id)
        if uplinks:
            uplink_path = self.output_dir / f"uplink_data_{tag}.csv"
            keys = list(uplinks[0].keys())
            with open(uplink_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=keys)
                writer.writeheader()
                writer.writerows(uplinks)
            created_files["csv_uplinks"] = uplink_path

        # 2. Downlinks CSV
        downlinks = self.db.fetch_all_downlinks(session_id=session_id)
        if downlinks:
            downlink_path = self.output_dir / f"downlink_data_{tag}.csv"
            keys = list(downlinks[0].keys())
            with open(downlink_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=keys)
                writer.writeheader()
                writer.writerows(downlinks)
            created_files["csv_downlinks"] = downlink_path

        # 3. RTT View CSV
        rtt_rows = self.db.fetch_rtt_view(session_id=session_id)
        if rtt_rows:
            rtt_path = self.output_dir / f"rtt_analysis_{tag}.csv"
            keys = list(rtt_rows[0].keys())
            with open(rtt_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=keys)
                writer.writeheader()
                writer.writerows(rtt_rows)
            created_files["csv_rtt"] = rtt_path

        # 4. Gateway Stats CSV
        gw_stats = self.db.fetch_gateway_stats(session_id=session_id)
        if gw_stats:
            gw_path = self.output_dir / f"gateway_stats_{tag}.csv"
            keys = list(gw_stats[0].keys())
            with open(gw_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=keys)
                writer.writeheader()
                writer.writerows(gw_stats)
            created_files["csv_gw_stats"] = gw_path

        # 5. PDR Summary CSV (Isolated specifically to this session)
        pdr_devices = self.db.get_pdr_statistics(session_id=session_id)
        if pdr_devices:
            pdr_path = self.output_dir / f"pdr_analysis_{tag}.csv"
            keys = list(pdr_devices[0].keys())
            with open(pdr_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=keys)
                writer.writeheader()
                writer.writerows(pdr_devices)
            created_files["csv_pdr"] = pdr_path

        return created_files

    def export_to_excel(self, tag: str, session_id: Optional[str] = None) -> Path:
        """Generates a styled multi-sheet Excel workbook with session isolation and history."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Style definitions
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        sub_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        regular_font = Font(name="Segoe UI", size=10)
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        highlight_font = Font(name="Segoe UI", size=11, bold=True, color="002060")

        def add_styled_sheet(sheet_title: str, rows: List[Dict], headers_map: Optional[Dict[str, str]] = None):
            if not rows:
                return
            ws = wb.create_sheet(title=sheet_title)
            keys = list(rows[0].keys())
            headers = [headers_map.get(k, k) if headers_map else k for k in keys]

            # Write header with style
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Fast row appending (O(N) streaming instead of O(N^2) cell-by-cell lookups)
            for row in rows:
                ws.append([row.get(k) for k in keys])

            # Auto-fit column widths efficiently (sample header + top 50 rows)
            sample_rows = rows[:50]
            for col_idx, k in enumerate(keys, start=1):
                header_text = str(headers[col_idx - 1])
                max_len = len(header_text)
                for r in sample_rows:
                    val_str = str(r.get(k) or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Sheet 1: Uplink Data (Filtered by session)
        uplinks = self.db.fetch_all_uplinks(session_id=session_id)
        add_styled_sheet("Uplink Data", uplinks)

        # Sheet 2: Downlink ACK (Filtered by session)
        downlinks = self.db.fetch_all_downlinks(session_id=session_id)
        add_styled_sheet("Downlink ACK", downlinks)

        # Sheet 3: Uplink-Downlink RTT (Filtered by session)
        rtt_rows = self.db.fetch_rtt_view(session_id=session_id)
        add_styled_sheet("Uplink-Downlink RTT", rtt_rows)

        # Sheet 4: Gateway Stats (Filtered by session)
        gw_stats = self.db.fetch_gateway_stats(session_id=session_id)
        add_styled_sheet("Gateway Stats", gw_stats)

        # Sheet 5: Summary Statistics (Session KPIs)
        summary = self.db.get_summary_statistics(session_id=session_id)
        ws_sum = wb.create_sheet(title="Summary Statistics")

        sess_label = session_id if session_id else "(Todas las sesiones)"
        ws_sum.append(["Métrica de Tesis (KPI)", f"Valor Sesión: {sess_label}", "Unidad / Descripción"])
        for col_idx in (1, 2, 3):
            cell = ws_sum.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        summary_rows = [
            # PDR KPI Section (Isolated)
            ("Packet Delivery Rate (PDR de Sesión)", f"{summary.get('global_pdr_percent', 0.0):.2f} %", "Porcentaje de paquetes entregados en esta ejecución"),
            ("Packet Loss Rate (PLR de Sesión)", f"{summary.get('global_plr_percent', 0.0):.2f} %", "Porcentaje de paquetes perdidos durante esta sesión"),
            ("Total Paquetes Transmitidos (Esperados)", summary.get("total_expected_packets", 0), "Calculado según secuencia de f_cnt en esta sesión"),
            ("Total Paquetes Uplink Recibidos", summary.get("total_uplinks", 0), "Paquetes capturados durante esta sesión"),
            ("Total Paquetes Perdidos Estimados", summary.get("total_lost_packets", 0), "Huecos en la secuencia de f_cnt de esta sesión"),
            # RF Quality Section
            ("RSSI Mínimo", f"{summary.get('min_rssi')} dBm" if summary.get('min_rssi') is not None else "N/A", "Potencia mínima recibida"),
            ("RSSI Promedio", f"{summary.get('avg_rssi', 0):.2f} dBm" if summary.get('avg_rssi') is not None else "N/A", "Potencia media del enlace"),
            ("RSSI Máximo", f"{summary.get('max_rssi')} dBm" if summary.get('max_rssi') is not None else "N/A", "Potencia máxima recibida"),
            ("SNR Mínimo", f"{summary.get('min_snr')} dB" if summary.get('min_snr') is not None else "N/A", "Relación señal a ruido mínima"),
            ("SNR Promedio", f"{summary.get('avg_snr', 0):.2f} dB" if summary.get('avg_snr') is not None else "N/A", "Relación señal a ruido media"),
            ("SNR Máximo", f"{summary.get('max_snr')} dB" if summary.get('max_snr') is not None else "N/A", "Relación señal a ruido máxima"),
            # Timing & Latency Section
            ("Airtime LoRa Promedio (T_aire)", f"{summary.get('avg_airtime', 0):.4f} s" if summary.get('avg_airtime') is not None else "N/A", "Tiempo de transmisión en el canal físico de radio"),
            ("Delay de Red y Servidor Promedio (T_red)", f"{summary.get('avg_delay', 0):.4f} s" if summary.get('avg_delay') is not None else "N/A", "Tránsito IP (Gateway -> TTN Server) + Cómputo NS/AS"),
            ("Latencia Total Subida E2E Promedio (T_E2E)", f"{summary.get('avg_total_uplink_delay', 0):.4f} s" if summary.get('avg_total_uplink_delay') is not None else "N/A", "Latencia total extremo a extremo: Airtime + Red/Servidor"),
            ("Muestras de RTT Confirmado", summary.get("total_rtt_samples", 0), "Paquetes confirmados con ACK"),
            ("Round Trip Time (RTT) Infraestructura Promedio", f"{summary.get('avg_rtt', 0):.4f} s" if summary.get('avg_rtt') is not None else "N/A", "Tiempo total Uplink RX → Downlink ACK Encolado"),
            ("RTT Físico Completo Promedio (rtt_phy_complete)", f"{summary.get('avg_rtt_phy_complete', 0):.4f} s" if summary.get('avg_rtt_phy_complete') is not None else "N/A", "airtime_s + rtt_s + airtime_downlink"),
            ("Espera Efectiva Ventana RX1 AU915 (effective_dl_window_delay)", f"{summary.get('avg_effective_dl_window_delay_au915', 0):.4f} s" if summary.get('avg_effective_dl_window_delay_au915') is not None else "N/A", "5000ms - rtt_s (Espera en buffer de Gateway hasta TX RF)"),
            ("Server Cloud Turnaround Promedio (T_nube)", f"{summary.get('avg_cloud_turnaround', 0):.4f} s" if summary.get('avg_cloud_turnaround') is not None else "N/A", "Tiempo interno en TTN Cloud: received_at -> gs_send_at"),
            ("Latencia de Scheduling NS Promedio", f"{summary.get('avg_ns_scheduling', 0):.4f} s" if summary.get('avg_ns_scheduling') is not None else "N/A", "Tiempo de programación del Network Server"),
            ("Delay de Transmisión GW ACK Promedio (T_bajada)", f"{summary.get('avg_tx_delay', 0):.4f} s" if summary.get('avg_tx_delay') is not None else "N/A", "Tiempo de despacho GS → Gateway y confirmación"),
        ]

        for metric, val, unit in summary_rows:
            ws_sum.append([metric, val, unit])
            row_idx = ws_sum.max_row
            ws_sum.cell(row=row_idx, column=1).font = bold_font
            ws_sum.cell(row=row_idx, column=2).font = highlight_font if "PDR" in metric else regular_font
            ws_sum.cell(row=row_idx, column=3).font = regular_font

        ws_sum.append([])
        ws_sum.append([])

        # Table 2: Per-Device PDR Breakdown for this session
        pdr_devs = summary.get("devices_pdr", [])
        if pdr_devs:
            sec_header_row = ws_sum.max_row + 1
            ws_sum.append(["Dispositivo (Device ID)", "Rangos FCnt (Segmentos)", "Reinicios", "Paquetes Esperados", "Paquetes Recibidos", "Paquetes Perdidos", "PDR (%)", "PLR (%)"])
            for col_idx in range(1, 9):
                cell = ws_sum.cell(row=sec_header_row, column=col_idx)
                cell.fill = sub_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for d in pdr_devs:
                ws_sum.append([
                    d["device_id"],
                    d.get("fcnt_ranges", f"{d.get('min_fcnt')}-{d.get('max_fcnt')}"),
                    d.get("resets_detected", 0),
                    d["expected_count"],
                    d["received_count"],
                    d["lost_count"],
                    f"{d['pdr_percent']:.2f} %",
                    f"{d['plr_percent']:.2f} %"
                ])
                cur_row = ws_sum.max_row
                ws_sum.cell(row=cur_row, column=1).font = bold_font
                for c in range(2, 9):
                    ws_sum.cell(row=cur_row, column=c).font = regular_font
                    ws_sum.cell(row=cur_row, column=c).alignment = Alignment(horizontal="center")

        for col in ws_sum.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # Sheet 6: Historial de Todas las Sesiones
        all_sessions = self.db.fetch_all_sessions()
        if all_sessions:
            ws_hist = wb.create_sheet(title="Historial de Sesiones")
            headers = ["ID de Sesión", "Inicio (UTC)", "Fin (UTC)", "Aplicación", "Gateway", "Modo", "Uplinks", "Downlinks", "PDR de Sesión (%)"]
            ws_hist.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws_hist.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for s in all_sessions:
                ws_hist.append([
                    s.get("session_id"),
                    s.get("started_at"),
                    s.get("ended_at") or "En curso",
                    s.get("app_id"),
                    s.get("gateway_id"),
                    s.get("collection_mode"),
                    s.get("uplinks_count", 0),
                    s.get("downlinks_count", 0),
                    f"{s.get('pdr_percent', 100.0):.2f} %"
                ])
                cur_row = ws_hist.max_row
                ws_hist.cell(row=cur_row, column=1).font = bold_font
                for c in range(2, len(headers) + 1):
                    ws_hist.cell(row=cur_row, column=c).font = regular_font

            for col in ws_hist.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_hist.column_dimensions[col_letter].width = max(max_len + 3, 14)

        excel_path = self.output_dir / f"lorawan_data_{tag}.xlsx"
        wb.save(excel_path)
        return excel_path
